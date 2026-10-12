"""
Сборка корпуса документов для RAG (модуль 5) из scripts/corpus/*.py → data/<category>/...

Часть документов сохраняется в других форматах, чтобы проверить multi-format ingestion (Б5.5):
PDF (pymupdf), DOCX (python-docx), HTML. Даты изменения файлов (mtime) разнесены по времени —
это поле created_at для фильтров «только свежие» (Б5.2).

    python -m scripts.build_corpus
"""

import html
import os
import re
from datetime import datetime
from pathlib import Path

from scripts.corpus import money, safety_app, taxes

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"

AS_PDF = {"taxes/social_deduction_medical.md", "taxes/property_deduction.md", "deposits_credits/deposit_insurance_asv.md",
          "deposits_credits/credit_psk.md", "fraud_security/stolen_card_actions.md", "app_faq/tariffs_v2.md",
          "budgeting/emergency_fund.md", "taxes/iis3.md"}
AS_DOCX = {"taxes/declaration_3ndfl.md", "deposits_credits/early_repayment.md", "budgeting/family_budget.md",
           "fraud_security/drops.md", "app_faq/import_statements.md", "taxes/self_employed_npd.md"}
AS_HTML = {"fraud_security/fake_bank_security_call.md", "budgeting/rule_50_30_20.md", "deposits_credits/key_rate.md",
           "app_faq/privacy.md", "taxes/children_standard_deduction.md", "fraud_security/phishing_sites.md"}
OLD = {"app_faq/tariffs_v1.md": datetime(2025, 11, 1), "deposits_credits/microloans_risk.md": datetime(2025, 6, 15),
       "budgeting/cashback.md": datetime(2025, 3, 10), "taxes/sport_deduction.md": datetime(2025, 9, 1)}
FRESH_DEFAULT = datetime(2026, 9, 20)


def lines(text: str) -> list[tuple[str, str]]:
    out = []
    for raw in text.strip().splitlines():
        if raw.startswith("# "):
            out.append(("h1", raw[2:]))
        elif raw.strip():
            out.append(("p", raw.strip()))
    return out


def to_pdf(text: str, path: Path) -> None:
    import pymupdf

    doc = pymupdf.open()
    story = "".join(f"<h1>{html.escape(t)}</h1>" if k == "h1" else f"<p>{html.escape(t)}</p>" for k, t in lines(text))
    writer = pymupdf.DocumentWriter(str(path))
    story_obj = pymupdf.Story(html=story, user_css="body{font-family:sans-serif;font-size:11pt} h1{font-size:16pt}")
    rect = pymupdf.paper_rect("a4") + (50, 50, -50, -50)
    more = True
    while more:
        device = writer.begin_page(pymupdf.paper_rect("a4"))
        more, _ = story_obj.place(rect)
        story_obj.draw(device)
        writer.end_page()
    writer.close()
    doc.close()


def to_docx(text: str, path: Path) -> None:
    import docx

    d = docx.Document()
    d.core_properties.author = "Команда Финпомощника"
    for kind, t in lines(text):
        d.add_heading(t, level=1) if kind == "h1" else d.add_paragraph(t)
    d.save(path)


def to_html(text: str, path: Path) -> None:
    body = "\n".join(f"<h1>{html.escape(t)}</h1>" if k == "h1" else f"<p>{html.escape(t)}</p>" for k, t in lines(text))
    title = re.sub(r"^# ", "", text.strip().splitlines()[0])
    path.write_text(f"<!doctype html><html lang=\"ru\"><head><meta charset=\"utf-8\"><title>{html.escape(title)}</title>"
                    f"</head><body>\n<nav>Финпомощник · база знаний</nav>\n<main>{body}</main>\n"
                    f"<footer>© Финпомощник</footer></body></html>", encoding="utf-8")


def main() -> None:
    docs = {**taxes.DOCS, **money.DOCS, **safety_app.DOCS}
    counts: dict[str, int] = {}
    for rel, text in docs.items():
        target = DATA / rel
        target.parent.mkdir(parents=True, exist_ok=True)
        if rel in AS_PDF:
            target = target.with_suffix(".pdf")
            to_pdf(text, target)
        elif rel in AS_DOCX:
            target = target.with_suffix(".docx")
            to_docx(text, target)
        elif rel in AS_HTML:
            target = target.with_suffix(".html")
            to_html(text, target)
        else:
            target.write_text(text.strip() + "\n", encoding="utf-8")
        ts = OLD.get(rel, FRESH_DEFAULT).timestamp()
        os.utime(target, (ts, ts))
        counts[target.suffix] = counts.get(target.suffix, 0) + 1
    print(f"Документов: {len(docs)}; по форматам: {counts}")


if __name__ == "__main__":
    main()
