"""
Б4.3 — вложения → content-part для Chat Completions.

  картинка  → {"type": "image_url", ...} с data-URI: мультимодальная модель видит исходные
              пиксели в том же вызове (отдельного Vision-вызова «картинка → текст» нет);
  голос     → текст → {"type": "text", "text": "[пользователь сказал голосом]: ..."}; два способа (SpeechToText):
              Whisper-1 при OPENAI_API_KEY, иначе аудио-модель через OpenRouter (input_audio, STT_MODEL).
              Оба принимают ogg/opus из Telegram напрямую — FFmpeg и subprocess не нужны;
  PDF/DOCX  → извлечённый текст (≤ 30 000 символов) → text-part.

Безопасность:
- тип файла определяется по сигнатуре байтов (magic bytes), а не по присланному content_type:
  иначе под видом картинки можно было бы протащить что угодно;
- лимиты размера, страниц PDF и распакованного размера DOCX (DOCX — это zip, защита от zip-бомб);
- текст документа — недоверенные данные: помечается как документ и проверяется на инъекции
  (см. ChatService), т. к. в PDF можно спрятать «игнорируй инструкции».
"""

import base64
import io
import zipfile
from dataclasses import dataclass
from typing import Literal

from fastapi import UploadFile
from openai import AsyncOpenAI

from app.services.usage_limits import record_usage

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_IMAGE_BYTES = 5 * 1024 * 1024  # картинка хранится base64 в истории — держим её небольшой
MAX_TEXT_CHARS = 30_000
PDF_MAX_PAGES = 50
DOCX_MAX_UNZIPPED = 50 * 1024 * 1024

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


VOICE_PREFIX = "[пользователь сказал голосом]:\n"
# формат input_audio по MIME из sniff_mime
AUDIO_FORMATS = {"audio/ogg": "ogg", "audio/mpeg": "mp3", "audio/mp4": "m4a", "audio/wav": "wav"}
TRANSCRIBE_PROMPT = ("Дословно расшифруй речь на аудио на языке оригинала. Верни только текст расшифровки, "
                     "без комментариев и кавычек. Если речи нет — верни пустую строку.")


@dataclass(frozen=True)
class SpeechToText:
    """whisper — OpenAI audio.transcriptions; chat — модель с аудио-входом через Chat Completions (OpenRouter)."""

    client: AsyncOpenAI
    kind: Literal["whisper", "chat"]
    model: str


class MediaError(ValueError):
    def __init__(self, message: str, status: int = 415):
        super().__init__(message)
        self.status = status


def sniff_mime(data: bytes) -> str | None:
    """Тип по сигнатуре содержимого."""
    if data.startswith(b"\xff\xd8\xff"):
        return "image/jpeg"
    if data.startswith(b"\x89PNG\r\n\x1a\n"):
        return "image/png"
    if data[:4] == b"RIFF" and data[8:12] == b"WEBP":
        return "image/webp"
    if data.startswith(b"%PDF-"):
        return "application/pdf"
    if data.startswith(b"OggS"):
        return "audio/ogg"
    if data.startswith(b"ID3") or data[:2] in (b"\xff\xfb", b"\xff\xf3", b"\xff\xf2"):
        return "audio/mpeg"
    if data[4:8] == b"ftyp":
        return "audio/mp4"  # m4a
    if data[:4] == b"RIFF" and data[8:12] == b"WAVE":
        return "audio/wav"
    if data.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                if "word/document.xml" in z.namelist():
                    return DOCX_MIME
        except zipfile.BadZipFile:
            return None
    return None


async def read_limited(media: UploadFile, limit: int = MAX_UPLOAD_BYTES) -> bytes:
    data = await media.read(limit + 1)
    if len(data) > limit:
        raise MediaError(f"Файл больше {limit // (1024 * 1024)} МБ", status=413)
    return data


async def media_to_part(media: UploadFile, stt: SpeechToText | None = None) -> dict:
    data = await read_limited(media)
    mime = sniff_mime(data)
    if mime is None:
        raise MediaError(f"Неподдерживаемый тип файла (заявлен {media.content_type!r})")

    if mime.startswith("image/"):
        if len(data) > MAX_IMAGE_BYTES:
            raise MediaError("Картинка больше 5 МБ", status=413)
        b64 = base64.b64encode(data).decode()
        return {"type": "image_url", "image_url": {"url": f"data:{mime};base64,{b64}"}}

    if mime.startswith("audio/"):
        transcript = await transcribe(data, mime, media.filename or "audio.ogg", stt)
        return {"type": "text", "text": VOICE_PREFIX + transcript}

    if mime == "application/pdf":
        return {"type": "text", "text": f"[документ PDF]:\n{extract_pdf_text(data)[:MAX_TEXT_CHARS]}"}

    if mime == DOCX_MIME:
        return {"type": "text", "text": f"[документ DOCX]:\n{extract_docx_text(data)[:MAX_TEXT_CHARS]}"}

    raise MediaError(f"Неподдерживаемый тип файла: {mime}")


async def transcribe(audio_bytes: bytes, mime: str, filename: str, stt: SpeechToText | None) -> str:
    if stt is None:
        raise MediaError("Распознавание голоса не настроено (нет ни OPENAI_API_KEY, ни ключа OpenRouter)", status=503)
    if stt.kind == "whisper":
        f = io.BytesIO(audio_bytes)
        f.name = filename  # SDK определяет формат по расширению
        result = await stt.client.audio.transcriptions.create(model=stt.model, file=f)
        text = result.text
    else:
        audio = {"data": base64.b64encode(audio_bytes).decode(), "format": AUDIO_FORMATS[mime]}
        response = await stt.client.chat.completions.create(model=stt.model, messages=[{"role": "user", "content": [
            {"type": "text", "text": TRANSCRIBE_PROMPT}, {"type": "input_audio", "input_audio": audio}]}])
        record_usage(response.usage, stt.model)  # распознавание тоже идёт в дневной лимит пользователя
        text = response.choices[0].message.content or ""
    text = text.strip()
    if not text:
        raise MediaError("Не удалось распознать речь — попробуйте записать ещё раз или напишите текстом", status=422)
    return text


def extract_pdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = reader.pages[:PDF_MAX_PAGES]
    text = "\n".join((p.extract_text() or "") for p in pages).strip()
    # Скан без текстового слоя: почти нет текста на многих страницах — честно сообщаем модели
    if len(pages) >= 5 and len(text) < 100:
        return "[похоже на скан без текстового слоя — текст извлечь не удалось]"
    if len(reader.pages) > PDF_MAX_PAGES:
        text += f"\n[обработаны первые {PDF_MAX_PAGES} страниц из {len(reader.pages)}]"
    return text or "[в документе нет извлекаемого текста]"


def extract_docx_text(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        if sum(i.file_size for i in z.infolist()) > DOCX_MAX_UNZIPPED:
            raise MediaError("DOCX слишком большой после распаковки", status=413)
    import docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts) or "[в документе нет текста]"
